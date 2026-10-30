"""Exemplos HTML editáveis derivados da geometria dos modelos PPTX; não altera os originais."""
import hashlib, html, json, os, pathlib
from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE
ROOT=pathlib.Path(__file__).resolve().parents[1];SK=ROOT/"skills/qp-brandbook";OUT=ROOT/"validacao/exemplos"
FONT_CSS="../../skills/qp-brandbook/assets/fonts.css"
def css_font(name):
    s=name or"Sora";weight=200 if"ExtraLight"in s else 300 if"Light"in s else 600 if"SemiBold"in s else 500 if"Medium"in s else 700 if"Bold"in s else 400
    family="Sora" if"Sora"in s else"Inter"if"Inter"in s else"Instrument Serif"if"Instrument"in s else s
    return family,weight
def rgb(color,default="#0D2240"):
    try:
        if color.type==MSO_COLOR_TYPE.RGB:return "#"+str(color.rgb)
    except (AttributeError,TypeError):pass
    return default
def text_style(run,pixel_scale=1):
    font=run.font;family,weight=css_font(font.name)
    return "font-family:'"+family+"';font-weight:"+str(700 if font.bold else weight)+";font-size:"+str(font.size.pt*96/72*pixel_scale if font.size else 24*pixel_scale)+"px;color:"+rgb(font.color)+";font-style:"+("italic"if font.italic else"normal")
def render_slide(prs,sl,texts,title):
    scale=960/prs.slide_width;parts=[];text_i=0
    try:bg=rgb(sl.background.fill.fore_color,"#FFFFFF")
    except:bg="#FFFFFF"
    for owner in [sl,sl.slide_layout,sl.slide_layout.slide_master]:
        blips=owner._element.xpath('./p:cSld/p:bg/p:bgPr/a:blipFill/a:blip')
        if blips:
            rid=blips[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed');part=owner.part.related_part(rid)
            name=hashlib.sha256(part.blob).hexdigest()[:20]+pathlib.Path(str(part.partname)).suffix
            (OUT/'assets'/name).write_bytes(part.blob)
            parts.append(f'<img alt="Fundo original do modelo" src="assets/{name}" style="position:absolute;inset:0;width:100%;height:100%">');break
    for shape in sl.shapes:
        x,y,w,h=[v*scale for v in [shape.left,shape.top,shape.width,shape.height]];style=f"position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;"
        if getattr(shape,"image",None):
            image=shape.image;name=hashlib.sha256(image.blob).hexdigest()[:20]+"."+image.ext;p=OUT/"assets"/name;p.parent.mkdir(exist_ok=True);p.write_bytes(image.blob)
            l,t,r,b=shape.crop_left,shape.crop_top,shape.crop_right,shape.crop_bottom
            parts.append(f'<div style="{style}overflow:hidden"><img alt="Elemento original do template" src="assets/{name}" style="position:absolute;left:{-l/(1-l-r)*100}%;top:{-t/(1-t-b)*100}%;width:{100/(1-l-r)}%;height:{100/(1-t-b)}%"></div>');continue
        if shape.has_text_frame:
            if not shape.text.strip():continue
            value=texts[text_i]if text_i<len(texts)else shape.text;text_i+=1
            tf=shape.text_frame;paragraph=tf.paragraphs[0];runs=paragraph.runs;run=runs[0]if runs else None;align="center"if str(paragraph.alignment).startswith("CENTER")else"right"if str(paragraph.alignment).startswith("RIGHT")else"left"
            ts=text_style(run,960/(prs.slide_width/914400*96))if run else"font:400 24px Sora";mt=tf.margin_top*scale;ml=tf.margin_left*scale;mr=tf.margin_right*scale
            parts.append(f'<div class="editable" contenteditable="true" spellcheck="false" style="{style}{ts};text-align:{align};padding:{mt}px {mr}px 0 {ml}px;line-height:1.25;white-space:pre-wrap;overflow:visible">{html.escape(value)}</div>');continue
        try:
            if shape.fill.type is not None:parts.append(f'<div style="{style}background:{rgb(shape.fill.fore_color,"transparent")};"></div>')
        except (AttributeError,TypeError):pass
    return f'<section class="slide" aria-label="{html.escape(title)}" style="height:{prs.slide_height*scale}px;background:{bg}">'+''.join(parts)+"</section>"
def document(title,body,note,lang="pt-BR"):
    return f'<!doctype html><html lang="{lang}"><meta charset="utf-8"><title>{html.escape(title)}</title><link rel="stylesheet" href="{FONT_CSS}"><style>*{{box-sizing:border-box}}body{{margin:0;background:#e7edf3;font-family:Inter,sans-serif;color:#0d2240}}header{{width:960px;padding:22px 0;margin:auto;font-size:14px;line-height:1.6}}.slide{{position:relative;width:960px;margin:0 auto 26px;overflow:hidden;background:white}}.editable:focus{{outline:2px solid #0066c6}}img{{display:block}}@media print{{body{{background:white}}header{{display:none}}.slide{{margin:0;break-after:page}}@page{{size:10in 5.625in;margin:0}}}}</style><header><strong>{html.escape(title)}</strong><br>{html.escape(note)}</header>{body}</html>'

def adapted_template(key,original,destination):
    """Título e apoio novos; região fotográfica vem do preview oficial recortado por CSS."""
    image=os.path.relpath(SK/original["path"],destination.parent)
    if key=="viajo-mucho":
        family,weight,height,lang="Lilita One",400,1280,"es-MX"
        body=f'''<section class="slide" aria-label="Viajo Mucho · adaptación" style="width:720px;height:{height}px;background:#0187E6">
<div class="editable" contenteditable="true" spellcheck="false" data-font="Lilita One" style="position:absolute;left:75px;top:165px;width:570px;height:155px;font:400 44px/1.12 'Lilita One';font-synthesis:none;color:white;text-align:center">TU PRÓXIMO VIAJE<br>EMPIEZA CONTIGO.</div>
<div style="position:absolute;top:350px;left:0;width:720px;height:690px;overflow:hidden"><img src="{image}" alt="Mosaico de destinos del modelo original" style="position:absolute;left:0;top:-350px;width:720px"></div>
<div class="editable" contenteditable="true" spellcheck="false" style="position:absolute;left:80px;top:1060px;width:560px;height:110px;font:400 27px/1.25 'Work Sans';color:white;text-align:center">Encuéntralo en<br><strong style="font-weight:600">viajomucho.com.mx</strong></div>
</section>'''
    else:
        family,weight,height,lang="Nunito",800,960,"pt-BR"
        body=f'''<section class="slide" aria-label="Rodon · adaptação" style="width:720px;height:{height}px;background:#0005C5">
<div style="position:absolute;inset:0 0 auto;height:610px;overflow:hidden"><img src="{image}" alt="Fotografia do modelo original Campos do Jordão" style="position:absolute;left:0;top:0;width:720px"></div>
<div class="editable" contenteditable="true" spellcheck="false" data-font="Nunito" style="position:absolute;left:55px;top:555px;width:610px;height:155px;padding:26px 30px;background:#0005C5;border-radius:32px;color:white;font:800 36px/1.15 Nunito;font-synthesis:none">Fim de semana em<br>Campos do Jordão</div>
<div class="editable" contenteditable="true" spellcheck="false" style="position:absolute;left:85px;top:740px;width:535px;height:90px;color:white;font:400 27px/1.3 Poppins">Roteiro completo<br>para sua viagem.</div>
<div aria-hidden="true" style="position:absolute;right:55px;bottom:38px;width:170px;height:65px;background:#FB8E00;border-radius:40px;color:white;font:400 54px/1 Inter;text-align:center">→</div>
</section>'''
    note="Adaptação autorizada: "+family+" substitui "+("Mauren"if key=="viajo-mucho"else"All Round Gothic")+". Título e apoio editáveis; imagens são recortes do preview do frame "+original["source_id"]+". Para editar fotografias ou ampliar, obtenha os arquivos de origem. Não altera o Figma."
    output=document(key+" · Google Fonts",body,note,lang)
    output=output.replace(FONT_CSS,os.path.relpath(SK/"assets/fonts.css",destination.parent))
    destination.parent.mkdir(parents=True,exist_ok=True);destination.write_text(output,encoding="utf8")
    return family,weight
def main():
    OUT.mkdir(parents=True,exist_ok=True);(OUT/"assets").mkdir(exist_ok=True)
    source=next((SK/"assets/templates/apresentacoes").glob("*modelo*"));prs=Presentation(source)
    selections=[(0,["Design","Clareza para criar." ]),(1,["Comece pelo público","Defina o objetivo, o canal e o formato.\nEscolha o modelo e os arquivos da marca.\nConfira a peça antes de compartilhar."]),(2,["Uma voz próxima","Clareza, cuidado e confiança."]),(3,["Cores com contexto","Escolha a paleta conforme a aplicação."])]
    deck="".join(render_slide(prs,prs.slides[i],texts,"Slide "+str(j+1))for j,(i,texts)in enumerate(selections))
    (OUT/"apresentacao-qp.html").write_text(document("Apresentação QP · exemplo local",deck,"Derivada do modelo oficial. Texto editável no navegador. Geometria importada do PPTX; validar diferenças de renderização antes de uso."),encoding="utf8")
    source_comm=next((SK/"assets/templates/recursos").glob("*comunicados*"));comm=Presentation(source_comm)
    text="Exemplo de comunicado interno.\n\nExplique a mudança e o próximo passo. Complete os detalhes com a área responsável antes de compartilhar."
    body=render_slide(comm,comm.slides[0],[text,"Um novo processo"],"Comunicado")
    (OUT/"comunicado-qp.html").write_text(document("Comunicado QP · exemplo local",body,"Modelo oficial preservado como referência. Não é anúncio de uma mudança real e não contém datas ou políticas inventadas."),encoding="utf8")
    index=json.loads((SK/"references/assets-index.json").read_text())["assets"]
    logo=next(a for a in index if a["provider"]=="figma"and a["category"]=="logos"and"COLOR"in a["name"]and"horizontal - color"in a["name"])
    photo=next(a for a in index if a["category"]=="fotografia"and a["path"].endswith(".png")and"jessy"in a["path"])
    lp="../../skills/qp-brandbook/"+logo["path"];pp="../../skills/qp-brandbook/"+photo["path"]
    post=f'<section class="slide" style="width:720px;height:900px;background:#0D2240"><div style="position:absolute;left:54px;top:54px;width:612px;height:500px;background:#0066C6;border-radius:30px;overflow:hidden"><img src="{pp}" alt="Recorte oficial do ensaio Jessy" style="width:100%;height:100%;object-fit:contain"></div><div class="editable" contenteditable="true" data-font="Bacalar Expanded" style="position:absolute;left:54px;top:610px;color:white;font:600 35px/1.2 \'Bacalar Expanded\';font-synthesis:none;width:612px;height:60px">AÇÃO QUE TE MOVE</div><div class="editable" contenteditable="true" style="position:absolute;left:54px;top:685px;width:600px;height:70px;color:#FFCC00;font:italic 400 42px/1.25 \'Instrument Serif\'">Sua próxima viagem.</div><img src="{lp}" alt="Quero Passagem" style="position:absolute;left:54px;bottom:40px;width:220px;filter:brightness(0) invert(1)"></section>'
    # Use an official negative logo instead of recoloring an image.
    neg=next(a for a in index if a["provider"]=="figma"and a["category"]=="logos"and"WHITE"in a["name"]and"qp - horizontal"in a["name"])
    post=post.replace(lp,"../../skills/qp-brandbook/"+neg["path"]).replace(";filter:brightness(0) invert(1)","")
    (OUT/"post-qp.html").write_text(document("Post QP · estudo com fundamentos",post,"Bacalar Expanded com acentos, do mesmo kit do Brand Navigator, peso 600 + Instrument Serif Italic. Estudo com paleta normativa e arquivos oficiais; não é reprodução exata do frame QP 1:1759."),encoding="utf8")
    cases=[{"file":"apresentacao-qp.html","source":"drive:1r80HtAMd2PTflVlzT-HXWC1xjIuEff-A3lP8uknESHo","source_slides":[i+1 for i,_ in selections],"status":"render-required"},{"file":"comunicado-qp.html","source":"drive:19lgUrEMkv7sYomWGXFMwfuqQQKP6mKlI3aBxlS94Fao","source_slides":[1],"status":"render-required"},{"file":"post-qp.html","source":"paginas/fundamentos/tipografia.md","assets":[neg["id"],photo["id"],"brand-navigator:bacalar-expanded-accented"],"required_font":"Bacalar Expanded","font_weight":600,"expected_postscript":"-","status":"render-required"}]
    for key,source_id in [("viajo-mucho","1:3470"),("rodon","42:42")]:
        original=next(a for a in index if a.get("source_id")==source_id and a["provider"]=="figma")
        family,weight=adapted_template(key,original,OUT/(key+".html"))
        adapted_template(key,original,SK/"assets/templates/adaptados"/(key+".html"))
        cases.append({"file":key+".html","source":original["id"],"status":"render-required","adaptation":"substituicao-tipografica-solicitada-2026-10-30","required_font":family,"font_weight":weight,"expected_postscript":"LilitaOne"if key=="viajo-mucho"else"Nunito-ExtraBold","image_limitation":"recorte-do-preview-oficial","commercial_claims":[]})
    (OUT/"casos.json").write_text(json.dumps(cases,ensure_ascii=False,indent=2)+"\n")
    print(json.dumps({"examples":len(cases),"adapted_templates":2,"fonts_pending":[]}))
if __name__=="__main__":main()
